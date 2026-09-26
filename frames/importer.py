from __future__ import annotations

from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook


PRODUCT_HEADERS = [
    "Категория",
    "Подкатегория",
    "Изображение",
    "Наименование товара",
    "Производство",
    "Описание",
    "Состав",
    "Цена",
]
SIZE_HEADERS = ["Размер"]
STOCK_HEADERS = [
    "Наименование товара",
    "Производство",
    "Размер",
    "Количество доступное для заказа",
]
USER_HEADERS = ["Фамилия", "Имя", "Отчетсво", "Логин", "Роль"]
ORDER_HEADERS = [
    "Номер заказа",
    "Дата заказа",
    "ФИО",
    "Категория",
    "Наименование товара",
    "Производство",
    "Размер",
    "Количество",
    "Цена за единицу",
]


class ImportValidationError(ValueError):
    pass


@dataclass(frozen=True)
class PreparedData:
    users: list[dict]
    products: list[dict]
    sizes: list[dict]
    stock_items: list[dict]
    orders: list[dict]
    order_items: list[dict]

    @property
    def counts(self) -> dict[str, int]:
        return {
            "users": len(self.users),
            "products": len(self.products),
            "sizes": len(self.sizes),
            "stock_items": len(self.stock_items),
            "orders": len(self.orders),
            "order_items": len(self.order_items),
        }


def normalize_text(value) -> str:
    return " ".join(str(value).replace("\xa0", " ").split())


def normalize_size(value) -> str:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return normalize_text(value)


def read_rows(path: Path, expected_headers: list[str]) -> list[dict]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)
        try:
            headers = [normalize_text(value) if value is not None else "" for value in next(iterator)]
        except StopIteration as error:
            raise ImportValidationError(f"Файл {path.name} пуст.") from error
        if headers != expected_headers:
            raise ImportValidationError(
                f"Неверные столбцы в {path.name}. Ожидались: {expected_headers}. "
                f"Получены: {headers}."
            )
        rows = []
        for row_number, values in enumerate(iterator, start=2):
            if all(value is None for value in values):
                continue
            row = dict(zip(headers, values))
            missing = [header for header in headers if row.get(header) is None]
            if missing:
                raise ImportValidationError(
                    f"В файле {path.name}, строка {row_number}, отсутствуют значения: {', '.join(missing)}."
                )
            rows.append(row)
        return rows
    finally:
        workbook.close()


def _resolve_product_id(name: str, manufacturer: str, products: list[dict]) -> int:
    normalized_name = normalize_text(name).casefold()
    normalized_manufacturer = normalize_text(manufacturer).casefold()
    manufacturer_matches = [
        product
        for product in products
        if product["manufacturer"].casefold() == normalized_manufacturer
    ]
    exact = [product for product in manufacturer_matches if product["name"].casefold() == normalized_name]
    if len(exact) == 1:
        return exact[0]["product_id"]
    prefix = [
        product
        for product in manufacturer_matches
        if product["name"].casefold().startswith(normalized_name)
        or normalized_name.startswith(product["name"].casefold())
    ]
    if len(prefix) == 1:
        return prefix[0]["product_id"]
    raise ImportValidationError(
        f"Не удалось однозначно сопоставить товар «{name}» производства «{manufacturer}»."
    )


def prepare_task2_data(resources_dir: Path) -> PreparedData:
    product_rows = read_rows(resources_dir / "Products_import.xlsx", PRODUCT_HEADERS)
    size_rows = read_rows(resources_dir / "Sizes_import.xlsx", SIZE_HEADERS)
    stock_rows = read_rows(resources_dir / "Stock_Items_import.xlsx", STOCK_HEADERS)
    user_rows = read_rows(resources_dir / "Users_import.xlsx", USER_HEADERS)
    order_rows = read_rows(resources_dir / "Orders_import.xlsx", ORDER_HEADERS)

    users = []
    login_keys = set()
    full_name_to_user_id = {}
    for user_id, row in enumerate(user_rows, start=1):
        login = normalize_text(row["Логин"])
        if login.casefold() in login_keys:
            raise ImportValidationError(f"Повторяющийся логин: {login}.")
        login_keys.add(login.casefold())
        user = {
            "user_id": user_id,
            "last_name": normalize_text(row["Фамилия"]),
            "first_name": normalize_text(row["Имя"]),
            "patronymic": normalize_text(row["Отчетсво"]),
            "login": login,
            "role": normalize_text(row["Роль"]),
        }
        full_name = normalize_text(
            f"{user['last_name']} {user['first_name']} {user['patronymic']}"
        )
        if full_name.casefold() in full_name_to_user_id:
            raise ImportValidationError(f"Повторяющееся ФИО пользователя: {full_name}.")
        full_name_to_user_id[full_name.casefold()] = user_id
        users.append(user)

    image_dir = resources_dir.parent / "images"
    products = []
    product_keys = set()
    for product_id, row in enumerate(product_rows, start=1):
        product = {
            "product_id": product_id,
            "category": normalize_text(row["Категория"]),
            "subcategory": normalize_text(row["Подкатегория"]),
            "image_filename": Path(normalize_text(row["Изображение"])).name,
            "name": normalize_text(row["Наименование товара"]),
            "manufacturer": normalize_text(row["Производство"]),
            "description": normalize_text(row["Описание"]),
            "composition": normalize_text(row["Состав"]),
            "price": Decimal(str(row["Цена"])).quantize(Decimal("0.01")),
        }
        key = (product["name"].casefold(), product["manufacturer"].casefold())
        if key in product_keys:
            raise ImportValidationError(
                f"Повторяющийся товар: {product['name']} | {product['manufacturer']}."
            )
        product_keys.add(key)
        if not (image_dir / product["image_filename"]).is_file():
            raise ImportValidationError(
                f"Изображение товара не найдено: {product['image_filename']}."
            )
        products.append(product)

    sizes = []
    size_to_id = {}
    for size_id, row in enumerate(size_rows, start=1):
        value = normalize_size(row["Размер"])
        if value in size_to_id:
            raise ImportValidationError(f"Повторяющийся размер: {value}.")
        size_to_id[value] = size_id
        sizes.append({"size_id": size_id, "value": value})

    stock_items = []
    stock_source_key_to_id = {}
    product_size_keys = set()
    for stock_item_id, row in enumerate(stock_rows, start=1):
        source_name = normalize_text(row["Наименование товара"])
        source_manufacturer = normalize_text(row["Производство"])
        size = normalize_size(row["Размер"])
        if size not in size_to_id:
            raise ImportValidationError(f"Неизвестный размер в остатках: {size}.")
        product_id = _resolve_product_id(source_name, source_manufacturer, products)
        unique_key = (product_id, size_to_id[size])
        if unique_key in product_size_keys:
            raise ImportValidationError(
                f"Повторяющийся остаток товара и размера: {source_name}, {size}."
            )
        product_size_keys.add(unique_key)
        quantity = int(row["Количество доступное для заказа"])
        if quantity < 0:
            raise ImportValidationError(f"Отрицательный остаток: {source_name}, размер {size}.")
        stock_items.append(
            {
                "stock_item_id": stock_item_id,
                "product_id": product_id,
                "size_id": size_to_id[size],
                "quantity": quantity,
            }
        )
        source_key = (
            source_name.casefold(),
            source_manufacturer.casefold(),
            size.casefold(),
        )
        stock_source_key_to_id[source_key] = stock_item_id

    grouped_order_rows = defaultdict(list)
    for row in order_rows:
        grouped_order_rows[int(row["Номер заказа"])].append(row)

    orders = []
    order_items = []
    order_item_id = 1
    for order_id in sorted(grouped_order_rows):
        rows = grouped_order_rows[order_id]
        dates = {_as_date(row["Дата заказа"]) for row in rows}
        clients = {normalize_text(row["ФИО"]) for row in rows}
        if len(dates) != 1 or len(clients) != 1:
            raise ImportValidationError(f"У заказа №{order_id} расходятся дата или клиент.")
        client_name = next(iter(clients))
        client_id = full_name_to_user_id.get(client_name.casefold())
        if client_id is None:
            raise ImportValidationError(f"Клиент заказа не найден в Users_import.xlsx: {client_name}.")
        orders.append(
            {
                "order_id": order_id,
                "order_date": next(iter(dates)),
                "client_id": client_id,
            }
        )
        combined = defaultdict(lambda: {"quantity": 0, "unit_price": None})
        for row in rows:
            source_key = (
                normalize_text(row["Наименование товара"]).casefold(),
                normalize_text(row["Производство"]).casefold(),
                normalize_size(row["Размер"]).casefold(),
            )
            stock_item_id = stock_source_key_to_id.get(source_key)
            if stock_item_id is None:
                raise ImportValidationError(
                    f"Позиция заказа №{order_id} отсутствует в Stock_Items_import.xlsx: {source_key}."
                )
            quantity = int(row["Количество"])
            if quantity <= 0:
                raise ImportValidationError(f"Некорректное количество в заказе №{order_id}.")
            unit_price = Decimal(str(row["Цена за единицу"])).quantize(Decimal("0.01"))
            current = combined[stock_item_id]
            if current["unit_price"] not in (None, unit_price):
                raise ImportValidationError(
                    f"У одной позиции заказа №{order_id} указаны разные цены."
                )
            current["quantity"] += quantity
            current["unit_price"] = unit_price
        for stock_item_id, item in combined.items():
            order_items.append(
                {
                    "order_item_id": order_item_id,
                    "order_id": order_id,
                    "stock_item_id": stock_item_id,
                    "quantity": item["quantity"],
                    "unit_price": item["unit_price"],
                }
            )
            order_item_id += 1

    prepared = PreparedData(users, products, sizes, stock_items, orders, order_items)
    expected_counts = {
        "users": 20,
        "products": 31,
        "sizes": 35,
        "stock_items": 92,
        "orders": 10,
        "order_items": 30,
    }
    if prepared.counts != expected_counts:
        raise ImportValidationError(
            f"Проверка количества записей не пройдена. Ожидалось: {expected_counts}. "
            f"Получено: {prepared.counts}."
        )
    return prepared


def _as_date(value) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for date_format in ("%d.%m.%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(normalize_text(value), date_format).date()
        except ValueError:
            continue
    raise ImportValidationError(f"Неверная дата: {value}.")


def sql_literal(value) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, date):
        return f"DATE '{value.isoformat()}'"
    if isinstance(value, (int, Decimal)):
        return str(value)
    return "'" + str(value).replace("'", "''") + "'"


def _insert_statement(table: str, columns: list[str], rows: list[dict]) -> str:
    values = []
    for row in rows:
        values.append(
            "(" + ", ".join(sql_literal(row[column]) for column in columns) + ")"
        )
    return (
        f"INSERT INTO {table} ({', '.join(columns)}) VALUES\n    "
        + ",\n    ".join(values)
        + ";"
    )


def generate_database_sql(prepared: PreparedData, schema_path: Path) -> str:
    blocks = [
        "-- База данных системы оформления заказов «Чудо Обувь»",
        "-- Сформировано по фактическим файлам TASK2.",
        "BEGIN;",
        schema_path.read_text(encoding="utf-8").strip(),
        _insert_statement(
            "users",
            ["user_id", "last_name", "first_name", "patronymic", "login", "role"],
            prepared.users,
        ),
        _insert_statement(
            "products",
            [
                "product_id",
                "category",
                "subcategory",
                "image_filename",
                "name",
                "manufacturer",
                "description",
                "composition",
                "price",
            ],
            prepared.products,
        ),
        _insert_statement("sizes", ["size_id", "value"], prepared.sizes),
        _insert_statement(
            "stock_items",
            ["stock_item_id", "product_id", "size_id", "quantity"],
            prepared.stock_items,
        ),
        _insert_statement(
            "orders",
            ["order_id", "order_date", "client_id"],
            prepared.orders,
        ),
        _insert_statement(
            "order_items",
            ["order_item_id", "order_id", "stock_item_id", "quantity", "unit_price"],
            prepared.order_items,
        ),
    ]
    for table, column in (
        ("users", "user_id"),
        ("products", "product_id"),
        ("sizes", "size_id"),
        ("stock_items", "stock_item_id"),
        ("orders", "order_id"),
        ("order_items", "order_item_id"),
    ):
        blocks.append(
            "SELECT setval(pg_get_serial_sequence(" 
            f"'{table}', '{column}'), (SELECT MAX({column}) FROM {table}), true);"
        )
    blocks.extend(
        [
            "COMMIT;",
            "",
            "-- Контрольные количества после импорта:",
            "-- users=20, products=31, sizes=35, stock_items=92, orders=10, order_items=30",
        ]
    )
    return "\n\n".join(blocks) + "\n"
